"""통합 테스트: 로컬 가짜 홍보사이트 → 수집 → 판별 → DB → 엑셀.

외부 인터넷에는 나가지 않습니다 (conftest 의 로컬 HTTP 서버만 사용).
"""

from __future__ import annotations

import pytest
from openpyxl import load_workbook

from illegal_site_bot.classifier import Classifier, load_rules
from illegal_site_bot.exporter import Exporter
from illegal_site_bot.fetcher import Fetcher
from illegal_site_bot.pipeline import Pipeline
from illegal_site_bot.storage import Storage


@pytest.fixture
def bot(config, promo_server):
    """가짜 홍보사이트 한 곳이 등록된 상태의 파이프라인 묶음."""
    storage = Storage(config.database_path)
    classifier = Classifier(load_rules(config.root / "config" / "rules.yaml"))
    fetcher = Fetcher(config.crawl)
    pipeline = Pipeline(
        config=config,
        storage=storage,
        classifier=classifier,
        fetcher=fetcher,
        renderer=None,
        pagination_patterns=("page=", "board"),
    )
    storage.upsert_source(promo_server, name="테스트 홍보사이트")
    try:
        yield {
            "config": config,
            "storage": storage,
            "classifier": classifier,
            "pipeline": pipeline,
            "source": promo_server,
        }
    finally:
        fetcher.close()
        storage.close()


def _collected_urls(storage: Storage) -> set[str]:
    return {row["url"] for row in storage.sites_for_export(0)}


class TestCycle:
    def test_cycle_reports_success(self, bot):
        stats = bot["pipeline"].run_cycle()
        assert stats.sources_total == 1
        assert stats.sources_ok == 1
        assert stats.sources_failed == 0
        assert stats.new_sites > 0

    def test_follows_internal_pagination(self, bot):
        stats = bot["pipeline"].run_cycle()
        # 메인 페이지 + 2페이지 (max_pages_per_source = 2)
        assert stats.pages_fetched == 2
        assert "https://baccarat-vip-365.club/" in _collected_urls(bot["storage"])

    def test_collects_links_from_every_extraction_path(self, bot):
        bot["pipeline"].run_cycle()
        urls = _collected_urls(bot["storage"])
        assert "https://casino-abc777.xyz/" in urls          # 이미지 배너 <a>
        assert "http://toto-safe-999.top/join" in urls       # 일반 <a>
        assert "https://slot-yamato-55.vip/" in urls         # data-href
        assert "https://holdem-king.cc/lobby" in urls        # onclick
        assert "https://freeya-dong.com/" in urls            # 본문 텍스트 abc[.]com
        assert "https://webtoon24-free.site/" in urls        # <script> 문자열

    def test_excludes_search_engines_and_low_score_noise(self, bot):
        bot["pipeline"].run_cycle()
        urls = _collected_urls(bot["storage"])
        assert not any("google.com" in url for url in urls)
        assert "https://plain-ad-frame.example/frame" not in urls

    def test_categorises_by_content(self, bot):
        bot["pipeline"].run_cycle()
        rows = {row["url"]: row for row in bot["storage"].sites_for_export(0)}
        assert rows["https://casino-abc777.xyz/"]["category"] == "gambling"
        assert rows["https://freeya-dong.com/"]["category"] == "adult"
        assert rows["https://t.me/promo_admin_contact"]["category"] == "contact"

    def test_records_which_promo_site_it_came_from(self, bot):
        bot["pipeline"].run_cycle()
        rows = {row["url"]: row for row in bot["storage"].sites_for_export(0)}
        site_id = rows["https://casino-abc777.xyz/"]["id"]
        assert bot["storage"].sources_for_site(site_id) == [bot["source"]]

    def test_source_result_is_recorded(self, bot):
        bot["pipeline"].run_cycle()
        row = bot["storage"].source_rows()[0]
        assert row["last_status"] == "200"
        assert row["consecutive_failures"] == 0
        assert row["found_total"] > 0
        assert row["last_crawled_at"]

    def test_second_cycle_finds_nothing_new(self, bot):
        first = bot["pipeline"].run_cycle()
        second = bot["pipeline"].run_cycle()
        assert second.new_sites == 0
        assert second.updated_sites == first.new_sites

        row = next(
            r for r in bot["storage"].sites_for_export(0)
            if r["url"] == "https://casino-abc777.xyz/"
        )
        assert row["seen_count"] == 2

    def test_stop_request_ends_cycle_early(self, bot):
        stats = bot["pipeline"].run_cycle(stop_check=lambda: True)
        assert stats.pages_fetched == 0
        assert stats.new_sites == 0

    def test_unreachable_source_is_marked_failed(self, bot, config):
        bot["storage"].upsert_source("http://127.0.0.1:9/", name="닫힌 포트")
        stats = bot["pipeline"].run_cycle()
        assert stats.sources_failed == 1
        failed = next(row for row in bot["storage"].source_rows() if ":9/" in row["url"])
        assert failed["consecutive_failures"] == 1
        assert failed["last_error"]

    def test_same_site_seen_from_two_promo_sources_scores_higher(self, bot, promo_server):
        # 점수가 100 으로 포화되지 않은 항목으로 확인합니다.
        target = "https://holdem-king.cc/lobby"
        bot["pipeline"].run_cycle()
        single = next(
            row for row in bot["storage"].sites_for_export(0) if row["url"] == target
        )
        assert single["score"] < 100

        # 같은 내용을 다른 주소로 한 번 더 등록하면 '여러 곳에서 홍보됨'이 됩니다.
        bot["storage"].upsert_source(promo_server + "index.html", name="두 번째 경로")
        bot["pipeline"].run_cycle()
        doubled = next(
            row for row in bot["storage"].sites_for_export(0) if row["url"] == target
        )
        assert doubled["distinct_sources"] == 2
        assert doubled["score"] > single["score"]


class TestExcelExport:
    @pytest.fixture
    def exported(self, bot):
        bot["pipeline"].run_cycle()
        exporter = Exporter(bot["config"], bot["storage"], bot["classifier"])
        return exporter.export()

    def test_creates_configured_file(self, exported, config):
        assert exported.path == config.export_path
        assert exported.path.is_file()
        assert exported.rows > 0

    def test_has_expected_sheets(self, exported):
        workbook = load_workbook(exported.path)
        assert workbook.sheetnames == [
            "요약",
            "불법사이트목록",
            "신규_최근24시간",
            "연락채널",
            "홍보사이트_수집원",
            "홍보사이트_후보",
            "실행이력",
        ]

    def test_main_sheet_lists_hosts_with_scores(self, exported):
        sheet = load_workbook(exported.path)["불법사이트목록"]
        headers = [cell.value for cell in sheet[1]]
        assert headers[:6] == ["번호", "호스트", "대표 URL", "URL 수", "발견된 URL", "도메인"]

        rows = {row[1]: row for row in sheet.iter_rows(min_row=2, values_only=True)}
        assert "casino-abc777.xyz" in rows
        casino = rows["casino-abc777.xyz"]
        assert casino[2] == "https://casino-abc777.xyz/"   # 대표 URL
        assert casino[5] == "casino-abc777.xyz"            # 도메인
        assert casino[6] == "도박/베팅"                     # 카테고리
        assert isinstance(casino[8], int) and casino[8] > 0  # 점수

    def test_urls_are_plain_text_by_default(self, exported):
        sheet = load_workbook(exported.path)["불법사이트목록"]
        # 실수로 접속하지 않도록 기본값은 하이퍼링크가 아닙니다.
        assert sheet.cell(row=2, column=3).hyperlink is None

    def test_contacts_go_to_their_own_sheet(self, exported):
        workbook = load_workbook(exported.path)
        contacts = [
            row[1] for row in workbook["연락채널"].iter_rows(min_row=2, values_only=True)
        ]
        main = [
            row[1] for row in workbook["불법사이트목록"].iter_rows(min_row=2, values_only=True)
        ]
        # 연락 채널은 묶지 않으므로 계정별 전체 URL 이 그대로 남습니다.
        assert "https://t.me/promo_admin_contact" in contacts
        assert "t.me" not in main

    def test_promo_source_sheet_lists_the_source(self, exported, promo_server):
        sheet = load_workbook(exported.path)["홍보사이트_수집원"]
        rows = [row for row in sheet.iter_rows(min_row=2, values_only=True)]
        assert any(row[1] == promo_server for row in rows)
        assert any(row[3] == "수동" for row in rows)   # 출처
        assert any(row[5] == "ON" for row in rows)     # 사용

    def test_new_sheet_contains_todays_finds(self, exported):
        sheet = load_workbook(exported.path)["신규_최근24시간"]
        hosts = [row[1] for row in sheet.iter_rows(min_row=2, values_only=True)]
        assert "casino-abc777.xyz" in hosts

    def test_export_is_rewritable(self, bot, exported):
        exporter = Exporter(bot["config"], bot["storage"], bot["classifier"])
        again = exporter.export()
        assert again.path == exported.path
        assert again.rows == exported.rows

    def test_marked_status_shows_up_in_excel(self, bot):
        bot["pipeline"].run_cycle()
        bot["storage"].set_site_status("https://casino-abc777.xyz/", "reported", "신고 완료")
        result = Exporter(bot["config"], bot["storage"], bot["classifier"]).export()

        sheet = load_workbook(result.path)["불법사이트목록"]
        row = next(
            row for row in sheet.iter_rows(min_row=2, values_only=True)
            if row[1] == "casino-abc777.xyz"
        )
        assert row[20] == "신고완료"
        assert row[21] == "신고 완료"


class TestRobots:
    def test_disallowed_source_is_skipped(self, bot, monkeypatch):
        """robots.txt 가 막으면 수집하지 않습니다."""
        fetcher = bot["pipeline"].fetcher
        assert fetcher.robots is not None
        monkeypatch.setattr(fetcher.robots, "allowed", lambda url: False)

        stats = bot["pipeline"].run_cycle()
        assert stats.sources_failed == 1
        assert stats.new_sites == 0


class TestDiscoveryLoop:
    """자가 확장 전체 흐름: 시드 수집 → 다른 홍보사이트 발견 → 승인 → 그곳까지 수집."""

    @pytest.fixture
    def loop(self, config, promo_server, partner_server):
        from illegal_site_bot.discovery import Discovery

        storage = Storage(config.database_path)
        classifier = Classifier(load_rules(config.root / "config" / "rules.yaml"))
        fetcher = Fetcher(config.crawl)
        discovery = Discovery(config, storage, classifier, fetcher)
        pipeline = Pipeline(
            config=config,
            storage=storage,
            classifier=classifier,
            fetcher=fetcher,
            renderer=None,
            pagination_patterns=("page=", "board"),
            discovery=discovery,
        )
        storage.upsert_source(promo_server, name="시드 홍보사이트")
        try:
            yield {
                "storage": storage,
                "pipeline": pipeline,
                "discovery": discovery,
                "seed": promo_server,
                "partner": partner_server,
                "config": config,
                "classifier": classifier,
            }
        finally:
            fetcher.close()
            storage.close()

    def test_cycle_discovers_the_linked_promo_site(self, loop):
        stats = loop["pipeline"].run_cycle()
        assert stats.candidates_added == 1

        rows = loop["storage"].sources_by_state("discovered")
        assert [row["url"] for row in rows] == [loop["partner"]]
        assert rows[0]["depth"] == 1
        assert rows[0]["origin"] == "auto"

    def test_discovered_site_is_not_crawled_before_approval(self, loop):
        loop["pipeline"].run_cycle()
        loop["discovery"].evaluate_pending()

        crawled = [source.url for source in loop["storage"].list_sources(enabled_only=True)]
        assert crawled == [loop["seed"]]

    def test_evaluation_queues_it_with_a_high_score(self, loop):
        loop["pipeline"].run_cycle()          # 시드에서 불법사이트들을 먼저 수집
        loop["discovery"].evaluate_pending()  # 그 결과가 A 신호로 쓰입니다

        row = loop["storage"].source_by_url(loop["partner"])
        assert row["state"] == "pending"
        assert row["promo_score"] >= loop["config"].discovery.auto_approve_score
        assert "이미 수집된 불법사이트" in row["promo_reasons"]

    def test_approved_site_is_crawled_and_yields_new_urls(self, loop):
        loop["pipeline"].run_cycle()
        loop["discovery"].evaluate_pending()
        assert loop["discovery"].approve(loop["partner"], by="테스트") is True

        # 승인 뒤 사이클에서는 그 사이트까지 돌면서 거기서만 보이는 URL 을 찾습니다.
        before = _collected_urls(loop["storage"])
        assert "https://newly-found-casino-42.top/" not in before

        stats = loop["pipeline"].run_cycle()
        assert stats.sources_total == 2
        after = _collected_urls(loop["storage"])
        assert "https://newly-found-casino-42.top/" in after

    def test_rejected_site_is_never_crawled(self, loop):
        loop["pipeline"].run_cycle()
        loop["discovery"].evaluate_pending()
        assert loop["discovery"].reject(loop["partner"], by="테스트") is True

        stats = loop["pipeline"].run_cycle()
        assert stats.sources_total == 1
        assert "https://newly-found-casino-42.top/" not in _collected_urls(loop["storage"])

    def test_discovery_can_be_turned_off(self, config, promo_server, partner_server):
        import dataclasses

        from illegal_site_bot.discovery import Discovery

        off = dataclasses.replace(
            config, discovery=dataclasses.replace(config.discovery, enabled=False)
        )
        storage = Storage(off.database_path)
        classifier = Classifier(load_rules(off.root / "config" / "rules.yaml"))
        fetcher = Fetcher(off.crawl)
        try:
            pipeline = Pipeline(
                config=off,
                storage=storage,
                classifier=classifier,
                fetcher=fetcher,
                pagination_patterns=("page=", "board"),
                discovery=Discovery(off, storage, classifier, fetcher),
            )
            storage.upsert_source(promo_server, name="시드")
            stats = pipeline.run_cycle()
            assert stats.candidates_added == 0
            assert storage.sources_by_state("discovered") == []
        finally:
            fetcher.close()
            storage.close()

    def test_candidate_sheet_shows_up_in_excel(self, loop):
        loop["pipeline"].run_cycle()
        loop["discovery"].evaluate_pending()
        result = Exporter(loop["config"], loop["storage"], loop["classifier"]).export()

        sheet = load_workbook(result.path)["홍보사이트_후보"]
        rows = [row for row in sheet.iter_rows(min_row=2, values_only=True)]
        candidate = next(row for row in rows if row[1] == loop["partner"])
        assert candidate[3] == "승인 대기"
        assert candidate[4] > 0                       # 점수
        assert candidate[5] == 1                      # 깊이
        assert "이미 수집된 불법사이트" in candidate[6]  # 판별 근거
        assert candidate[7] == loop["seed"]           # 발견 경로


class TestCsvExport:
    """외부 시스템(AI 재분류)이 읽어갈 CSV."""

    @pytest.fixture
    def exported(self, bot):
        bot["pipeline"].run_cycle()
        return Exporter(bot["config"], bot["storage"], bot["classifier"]).export()

    def _rows(self, path):
        import csv

        with open(path, encoding="utf-8-sig", newline="") as handle:
            return list(csv.DictReader(handle))

    def test_creates_both_csv_files(self, exported, config):
        assert exported.csv_path == config.export_dir / "urls.csv"
        assert exported.csv_new_path == config.export_dir / "urls_new.csv"
        assert exported.csv_path.is_file() and exported.csv_new_path.is_file()

    def test_has_the_requested_columns_with_url_first(self, exported):
        rows = self._rows(exported.csv_path)
        columns = list(rows[0].keys())
        assert columns[:4] == ["url", "score", "first_seen_at", "last_seen_at"]
        for required in ("host", "domain", "category", "alive", "promo_sites"):
            assert required in columns

    def test_starts_with_utf8_bom_so_excel_opens_it(self, exported):
        # BOM 이 없으면 한글 엑셀에서 글자가 깨집니다.
        assert exported.csv_path.read_bytes()[:3] == b"\xef\xbb\xbf"

    def test_rows_are_per_url_not_grouped(self, exported):
        """엑셀은 호스트 단위로 묶지만 CSV 는 원본 URL 을 그대로 넘깁니다."""
        urls = [row["url"] for row in self._rows(exported.csv_path)]
        assert "http://toto-safe-999.top/join" in urls
        assert "https://holdem-king.cc/lobby" in urls

    def test_score_and_timestamps_are_filled(self, exported):
        row = next(
            row for row in self._rows(exported.csv_path)
            if row["url"] == "https://casino-abc777.xyz/"
        )
        assert int(row["score"]) > 0
        assert row["first_seen_at"] and row["last_seen_at"]
        assert row["category"] == "gambling"
        assert row["promo_sites"]

    def test_contacts_are_not_in_the_feed(self, exported):
        urls = [row["url"] for row in self._rows(exported.csv_path)]
        assert not any("t.me" in url for url in urls)

    def test_reclassified_sites_are_excluded(self, bot):
        bot["pipeline"].run_cycle()
        bot["storage"].set_site_status("https://casino-abc777.xyz/", "ignored", "오탐")
        result = Exporter(bot["config"], bot["storage"], bot["classifier"]).export()
        urls = [row["url"] for row in self._rows(result.csv_path)]
        assert "https://casino-abc777.xyz/" not in urls

    def test_new_file_holds_only_recent_finds(self, exported):
        full = self._rows(exported.csv_path)
        recent = self._rows(exported.csv_new_path)
        assert 0 < len(recent) <= len(full)

    def test_min_score_filter(self, bot):
        import dataclasses

        strict = dataclasses.replace(
            bot["config"],
            export=dataclasses.replace(bot["config"].export, csv_min_score=90),
        )
        bot["pipeline"].run_cycle()
        result = Exporter(strict, bot["storage"], bot["classifier"]).export()
        scores = [int(row["score"]) for row in self._rows(result.csv_path)]
        assert scores and all(score >= 90 for score in scores)

    def test_can_be_turned_off(self, bot):
        import dataclasses

        off = dataclasses.replace(
            bot["config"],
            export=dataclasses.replace(bot["config"].export, csv_enabled=False),
        )
        bot["pipeline"].run_cycle()
        result = Exporter(off, bot["storage"], bot["classifier"]).export()
        assert result.csv_path is None
        assert not (bot["config"].export_dir / "urls.csv").exists()


class TestExportRegressions:
    """코드 리뷰에서 나온 출력 관련 버그들의 재발 방지."""

    def test_group_bonus_can_lift_a_site_over_min_score(self, config, promo_server):
        """점수 필터를 묶기 전에 걸면, 묶어서 붙는 가산점이 무의미해집니다."""
        import dataclasses

        storage = Storage(config.database_path)
        classifier = Classifier(load_rules(config.root / "config" / "rules.yaml"))
        try:
            # 각각 26점(기준 30 미달)이지만 서로 다른 홍보사이트에서 발견된 같은 호스트.
            for path, promo in (("/", "https://promo1.test/"), ("/join", "https://promo2.test/")):
                site_id, _ = storage.record_site(
                    f"https://target.com{path}",
                    category="gambling",
                    category_label="도박/베팅",
                    base_score=26,
                )
                storage.record_observation(
                    site_id, None, promo, promo, "", "a_href", False, 26
                )
                storage.refresh_site_score(site_id, 6, 24)

            assert [row["score"] for row in storage.sites_for_export(0)] == [26, 26]

            strict = dataclasses.replace(
                config, export=dataclasses.replace(config.export, min_score=30)
            )
            result = Exporter(strict, storage, classifier).export()
            sheet = load_workbook(result.path)["불법사이트목록"]
            rows = [row for row in sheet.iter_rows(min_row=2, values_only=True)]

            assert [row[1] for row in rows] == ["target.com"]
            assert rows[0][8] == 32          # 26 + 홍보사이트 2곳 가산
            assert rows[0][12] == 2          # 홍보사이트 수
        finally:
            storage.close()

    def test_csv_timestamps_are_utc_iso(self, bot):
        """기계가 읽는 피드라 시간대를 알 수 있어야 합니다."""
        import csv
        import re

        bot["pipeline"].run_cycle()
        result = Exporter(bot["config"], bot["storage"], bot["classifier"]).export()
        with open(result.csv_path, encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle))

        pattern = re.compile(r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}Z$")
        assert rows, "CSV 가 비어 있습니다"
        for row in rows:
            assert pattern.match(row["first_seen_at"]), row["first_seen_at"]
            assert pattern.match(row["last_seen_at"]), row["last_seen_at"]

    def test_control_characters_do_not_break_the_export(self, bot):
        """제어문자가 든 값이 DB 에 남아 있어도 엑셀 저장은 계속돼야 합니다."""
        bot["pipeline"].run_cycle()
        with bot["storage"]._lock:
            bot["storage"]._conn.execute(
                "UPDATE sites SET title = ?, memo = ? WHERE url = ?",
                ("나쁜\x07제목", "메모\x01", "https://casino-abc777.xyz/"),
            )
            bot["storage"]._conn.commit()

        result = Exporter(bot["config"], bot["storage"], bot["classifier"]).export()
        assert result.path.is_file()


class TestLanguageFilterInPipeline:
    """수집 사이클에서 외국어 사이트가 DB 에 저장되지 않는지."""

    @pytest.fixture
    def bot_with_language(self, config, promo_server, partner_server):
        from illegal_site_bot.language import LanguageDetector

        storage = Storage(config.database_path)
        classifier = Classifier(load_rules(config.root / "config" / "rules.yaml"))
        fetcher = Fetcher(config.crawl)
        language = LanguageDetector(
            config.language_filter, fetcher=fetcher, storage=storage
        )
        pipeline = Pipeline(
            config=config,
            storage=storage,
            classifier=classifier,
            fetcher=fetcher,
            pagination_patterns=("page=", "board"),
            language=language,
        )
        storage.upsert_source(promo_server, name="시드")
        try:
            yield {
                "config": config,
                "storage": storage,
                "classifier": classifier,
                "pipeline": pipeline,
                "language": language,
                "partner": partner_server,
            }
        finally:
            fetcher.close()
            storage.close()

    def _seed_pages(self, bot, korean: bool) -> None:
        """수집 대상에 한국어/영어 페이지를 직접 등록해 사이클을 돌립니다."""
        page = "korean" if korean else "english"
        bot["storage"].upsert_source(bot["partner"] + page, name=page)

    def test_english_page_is_judged_foreign(self, config, bot_with_language):
        """로컬 픽스처의 영어 페이지를 실제로 받아 판별합니다."""
        import dataclasses

        from illegal_site_bot.language import LanguageDetector

        bot = bot_with_language
        settings = dataclasses.replace(
            config.language_filter, fetch_when_unknown=True
        )
        detector = LanguageDetector(
            settings, fetcher=bot["pipeline"].fetcher, storage=bot["storage"]
        )
        detector.begin_cycle()
        allowed, verdict = detector.allows(bot_with_language["partner"] + "english")
        assert allowed is False and verdict.language == "foreign"

    def test_korean_sites_are_still_collected(self, bot_with_language):
        bot = bot_with_language
        bot["pipeline"].run_cycle()
        urls = {row["url"] for row in bot["storage"].sites_for_export(0)}
        # 기존 픽스처의 도박 사이트들은 .xyz/.top 이라 접속 판별 대상이지만,
        # 접속되지 않으므로 '판별 불가 → 보존' 규칙에 따라 남아야 합니다.
        assert "https://casino-abc777.xyz/" in urls

    def test_contact_channels_bypass_the_filter(self, bot_with_language):
        """t.me 는 언어가 없는 대상이라 필터에 걸리면 안 됩니다."""
        bot = bot_with_language
        bot["pipeline"].run_cycle()
        urls = {row["url"] for row in bot["storage"].sites_for_export(0)}
        assert "https://t.me/promo_admin_contact" in urls

    def test_dropped_count_is_reported(self, bot_with_language):
        bot = bot_with_language
        # 영어 사이트를 확실히 후보로 만들기 위해 캐시에 심어둡니다.
        bot["storage"].record_domain_language(
            "casino-abc777.xyz", language="foreign", reason="테스트"
        )
        stats = bot["pipeline"].run_cycle()
        assert stats.dropped_foreign >= 1

        urls = {row["url"] for row in bot["storage"].sites_for_export(0)}
        assert "https://casino-abc777.xyz/" not in urls

    def test_filter_off_keeps_everything(self, config, promo_server):
        import dataclasses

        from illegal_site_bot.language import LanguageDetector

        off = dataclasses.replace(
            config,
            language_filter=dataclasses.replace(config.language_filter, enabled=False),
        )
        storage = Storage(off.database_path)
        classifier = Classifier(load_rules(off.root / "config" / "rules.yaml"))
        fetcher = Fetcher(off.crawl)
        try:
            storage.record_domain_language(
                "casino-abc777.xyz", language="foreign", reason="테스트"
            )
            pipeline = Pipeline(
                config=off,
                storage=storage,
                classifier=classifier,
                fetcher=fetcher,
                pagination_patterns=("page=", "board"),
                language=LanguageDetector(
                    off.language_filter, fetcher=fetcher, storage=storage
                ),
            )
            storage.upsert_source(promo_server, name="시드")
            stats = pipeline.run_cycle()

            assert stats.dropped_foreign == 0
            urls = {row["url"] for row in storage.sites_for_export(0)}
            assert "https://casino-abc777.xyz/" in urls
        finally:
            fetcher.close()
            storage.close()
