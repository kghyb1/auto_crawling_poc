"""수집 결과를 엑셀(.xlsx) 로 정리합니다.

엑셀은 매번 새로 만드는 '보고서'입니다. 사람이 엑셀에 직접 적은 메모는
다음 출력 때 사라지므로, 처리 상태나 메모는 ``python bot.py mark`` 로
DB 에 남겨야 합니다(그러면 엑셀에도 계속 따라옵니다).

시트 구성
    요약              한눈에 보는 수치
    불법사이트목록     수집된 URL 전체 (핵심 시트)
    신규_최근24시간    지난 24시간에 처음 발견된 것
    연락채널          텔레그램/카톡 등 접촉 채널
    홍보사이트_수집원  어디를 돌고 있는지 + 마지막 결과
    홍보사이트_후보    봇이 새로 찾아낸 홍보사이트 (승인 대기/기각)
    실행이력          사이클별 통계
"""

from __future__ import annotations

import csv
import logging
import os
import re
import shutil
import sqlite3
import tempfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Sequence

from openpyxl import Workbook
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .aggregation import (
    DEFAULT_GROUP_MODE,
    GROUP_LABELS,
    SiteGroup,
    group_sites,
)
from .classifier import CONTACT_CATEGORY, Classifier
from .config import Config
from .storage import Storage
from .timeutil import days_ago, humanize_duration, local_str, now_utc

log = logging.getLogger(__name__)

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
TITLE_FONT = Font(bold=True, size=14)
LABEL_FONT = Font(bold=True)
THIN_BORDER = Border(*(Side(style="thin", color="D9D9D9"),) * 4)

RISK_FILLS = {
    "높음": PatternFill("solid", fgColor="FFC7CE"),
    "보통": PatternFill("solid", fgColor="FFEB9C"),
    "낮음": PatternFill("solid", fgColor="E2EFDA"),
}

ALIVE_TEXT = {1: "생존", 0: "접속불가", None: "미확인"}
STATUS_TEXT = {
    "new": "신규",
    "confirmed": "확인",
    "reported": "신고완료",
    "ignored": "제외",
}

# (헤더, 너비) — 두 번째 컬럼 이름은 묶음 기준에 따라 달라집니다.
SITE_COLUMNS: tuple[tuple[str, int], ...] = (
    ("번호", 6),
    ("사이트", 34),
    ("대표 URL", 44),
    ("URL 수", 8),
    ("발견된 URL", 46),
    ("도메인", 24),
    ("카테고리", 12),
    ("위험도", 8),
    ("점수", 7),
    ("최초 발견", 18),
    ("최근 발견", 18),
    ("발견 횟수", 9),
    ("홍보사이트 수", 12),
    ("발견된 홍보사이트", 38),
    ("일치 키워드", 26),
    ("판별 근거", 26),
    ("경유 URL", 26),
    ("생존", 9),
    ("HTTP", 7),
    ("확인 시각", 18),
    ("처리 상태", 10),
    ("메모", 22),
)


def site_columns(mode: str) -> tuple[tuple[str, int], ...]:
    """묶음 기준에 맞춰 두 번째 컬럼 이름을 바꿔 돌려줍니다."""
    columns = list(SITE_COLUMNS)
    columns[1] = (GROUP_LABELS.get(mode, "사이트"), columns[1][1])
    return tuple(columns)


SOURCE_COLUMNS: tuple[tuple[str, int], ...] = (
    ("번호", 6),
    ("홍보사이트 URL", 50),
    ("이름", 22),
    ("출처", 7),
    ("깊이", 6),
    ("사용", 7),
    ("렌더링", 8),
    ("최대 페이지", 11),
    ("마지막 수집", 18),
    ("마지막 상태", 12),
    ("연속 실패", 10),
    ("누적 수집", 10),
    ("마지막 오류", 40),
    ("비고", 24),
)

CANDIDATE_COLUMNS: tuple[tuple[str, int], ...] = (
    ("번호", 6),
    ("후보 URL", 50),
    ("사이트 제목", 28),
    ("상태", 12),
    ("점수", 7),
    ("깊이", 6),
    ("판별 근거", 52),
    ("발견 경로", 40),
    ("발견 시각", 18),
    ("평가 시각", 18),
)

RUN_COLUMNS: tuple[tuple[str, int], ...] = (
    ("시작", 18),
    ("종료", 18),
    ("소요", 12),
    ("대상", 7),
    ("성공", 7),
    ("실패", 7),
    ("페이지", 8),
    ("후보", 8),
    ("신규", 8),
    ("갱신", 8),
    ("신규후보", 9),
    ("자동승인", 9),
    ("외국어제외", 11),
    ("비고", 40),
)


@dataclass
class ExportResult:
    path: Path
    snapshot: Path | None
    rows: int
    new_rows: int
    contacts: int
    url_rows: int = 0
    csv_path: Path | None = None
    csv_new_path: Path | None = None
    csv_rows: int = 0
    warning: str = ""


#: 엑셀이 거부하는 제어문자. 추출 단계에서도 걸러내지만, 이전 버전이 저장해 둔
#: 값이 DB 에 남아 있을 수 있어 기록 직전에 한 번 더 막습니다. 여기서 새면
#: 엑셀 저장이 통째로 실패하고 그 뒤로 보고서가 갱신되지 않습니다.
_ILLEGAL_XLSX_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _safe(value: Any) -> Any:
    """엑셀 셀에 넣어도 안전한 값으로 만듭니다."""
    if isinstance(value, str):
        return _ILLEGAL_XLSX_CHARS.sub("", value)
    return value


def _style_header(sheet: Worksheet, columns: Sequence[tuple[str, int]], row: int = 1) -> None:
    for index, (title, width) in enumerate(columns, start=1):
        cell = sheet.cell(row=row, column=index, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[row].height = 24
    sheet.freeze_panes = sheet.cell(row=row + 1, column=1)
    last_column = get_column_letter(len(columns))
    sheet.auto_filter.ref = f"A{row}:{last_column}{row}"


def _write_url(sheet: Worksheet, row: int, column: int, url: str, clickable: bool) -> None:
    cell = sheet.cell(row=row, column=column, value=url)
    cell.alignment = Alignment(vertical="center")
    if clickable and url:
        cell.hyperlink = url
        cell.font = Font(color="0563C1", underline="single")
    else:
        # 기본은 일반 텍스트입니다. 실수로 불법사이트에 접속하는 것을 막기 위함입니다.
        cell.font = Font(name="Consolas", size=10)


class Exporter:
    def __init__(self, config: Config, storage: Storage, classifier: Classifier) -> None:
        self.config = config
        self.storage = storage
        self.classifier = classifier

    # -- 시트별 작성 -------------------------------------------------------
    @property
    def group_mode(self) -> str:
        return self.config.export.group_by or DEFAULT_GROUP_MODE

    def _group(
        self,
        rows: Sequence[sqlite3.Row],
        source_map: dict[int, list[str]],
        mode: str | None = None,
    ) -> list[SiteGroup]:
        rules = self.classifier.rules
        return group_sites(
            rows,
            source_map,
            mode=mode or self.group_mode,
            bonus_per_source=rules.cross_source_bonus_per_source,
            bonus_max=rules.cross_source_bonus_max,
        )

    def _write_sites_sheet(
        self, sheet: Worksheet, groups: Sequence[SiteGroup], mode: str | None = None
    ) -> None:
        _style_header(sheet, site_columns(mode or self.group_mode))
        clickable = self.config.export.clickable_links
        wrap = Alignment(wrap_text=True, vertical="top")
        center = Alignment(horizontal="center")

        for index, group in enumerate(groups, start=1):
            excel_row = index + 1
            risk = self.classifier.risk_label(group.score)

            sheet.cell(row=excel_row, column=1, value=index).alignment = center
            sheet.cell(row=excel_row, column=2, value=_safe(group.key)).font = Font(
                name="Consolas", size=10
            )
            _write_url(sheet, excel_row, 3, group.representative_url, clickable)

            count_cell = sheet.cell(row=excel_row, column=4, value=_safe(group.url_count))
            count_cell.alignment = center
            sheet.cell(row=excel_row, column=5, value=_safe(group.urls_text)).alignment = wrap
            sheet.cell(row=excel_row, column=6, value=_safe(group.domain))
            sheet.cell(row=excel_row, column=7, value=_safe(group.category_label))

            risk_cell = sheet.cell(row=excel_row, column=8, value=risk)
            risk_cell.alignment = center
            if risk in RISK_FILLS:
                risk_cell.fill = RISK_FILLS[risk]

            sheet.cell(row=excel_row, column=9, value=_safe(group.score)).alignment = center
            sheet.cell(row=excel_row, column=10, value=_safe(local_str(group.first_seen_at)))
            sheet.cell(row=excel_row, column=11, value=_safe(local_str(group.last_seen_at)))
            sheet.cell(row=excel_row, column=12, value=_safe(group.seen_count)).alignment = center
            sheet.cell(row=excel_row, column=13, value=_safe(group.distinct_sources)).alignment = center
            sheet.cell(row=excel_row, column=14, value=_safe(group.sources_text)).alignment = wrap
            sheet.cell(row=excel_row, column=15, value=_safe(group.matched_keywords)).alignment = wrap
            sheet.cell(row=excel_row, column=16, value=_safe(group.reasons)).alignment = wrap
            sheet.cell(row=excel_row, column=17, value=_safe(group.redirect_from))

            sheet.cell(
                row=excel_row, column=18, value=ALIVE_TEXT.get(group.alive, "미확인")
            ).alignment = center
            sheet.cell(row=excel_row, column=19, value=_safe(group.http_status)).alignment = center
            sheet.cell(row=excel_row, column=20, value=_safe(local_str(group.last_checked_at)))
            sheet.cell(
                row=excel_row, column=21, value=STATUS_TEXT.get(group.status, group.status)
            ).alignment = center
            sheet.cell(row=excel_row, column=22, value=_safe(group.memo)).alignment = wrap

        if groups:
            last = len(groups) + 1
            sheet.conditional_formatting.add(
                f"I2:I{last}",
                DataBarRule(
                    start_type="num",
                    start_value=0,
                    end_type="num",
                    end_value=100,
                    color="638EC6",
                    showValue=True,
                ),
            )

    def _write_sources_sheet(self, sheet: Worksheet, rows: Sequence[sqlite3.Row]) -> None:
        _style_header(sheet, SOURCE_COLUMNS)
        for index, row in enumerate(rows, start=1):
            excel_row = index + 1
            sheet.cell(row=excel_row, column=1, value=index).alignment = Alignment(
                horizontal="center"
            )
            _write_url(sheet, excel_row, 2, row["url"], False)
            sheet.cell(row=excel_row, column=3, value=_safe(row["name"]))
            sheet.cell(
                row=excel_row, column=4, value="자동" if row["origin"] == "auto" else "수동"
            ).alignment = Alignment(horizontal="center")
            sheet.cell(row=excel_row, column=5, value=int(row["depth"] or 0)).alignment = (
                Alignment(horizontal="center")
            )
            sheet.cell(
                row=excel_row, column=6, value="ON" if row["enabled"] else "OFF"
            ).alignment = Alignment(horizontal="center")
            sheet.cell(
                row=excel_row, column=7, value="예" if row["render"] else "-"
            ).alignment = Alignment(horizontal="center")
            sheet.cell(row=excel_row, column=8, value=_safe(row["max_pages"])).alignment = Alignment(
                horizontal="center"
            )
            sheet.cell(row=excel_row, column=9, value=_safe(local_str(row["last_crawled_at"])))
            sheet.cell(row=excel_row, column=10, value=_safe(row["last_status"]))
            failures = int(row["consecutive_failures"])
            failure_cell = sheet.cell(row=excel_row, column=11, value=failures)
            failure_cell.alignment = Alignment(horizontal="center")
            if failures >= 3:
                failure_cell.fill = RISK_FILLS["높음"]
            sheet.cell(row=excel_row, column=12, value=int(row["found_total"]))
            sheet.cell(row=excel_row, column=13, value=_safe(row["last_error"])).alignment = Alignment(
                wrap_text=True, vertical="top"
            )
            sheet.cell(row=excel_row, column=14, value=_safe(row["note"]))

    def _write_candidates_sheet(
        self, sheet: Worksheet, rows: Sequence[sqlite3.Row], origin_map: dict[int, str]
    ) -> None:
        _style_header(sheet, CANDIDATE_COLUMNS)
        state_text = {
            "discovered": "평가 대기",
            "pending": "승인 대기",
            "rejected": "기각",
            "auto_disabled": "자동 중지",
        }
        for index, row in enumerate(rows, start=1):
            excel_row = index + 1
            sheet.cell(row=excel_row, column=1, value=index).alignment = Alignment(
                horizontal="center"
            )
            _write_url(sheet, excel_row, 2, row["url"], False)
            sheet.cell(row=excel_row, column=3, value=_safe(row["name"]))

            state_cell = sheet.cell(
                row=excel_row, column=4, value=state_text.get(row["state"], row["state"])
            )
            state_cell.alignment = Alignment(horizontal="center")
            if row["state"] == "pending":
                state_cell.fill = RISK_FILLS["보통"]

            sheet.cell(row=excel_row, column=5, value=int(row["promo_score"] or 0)).alignment = (
                Alignment(horizontal="center")
            )
            sheet.cell(row=excel_row, column=6, value=int(row["depth"] or 0)).alignment = (
                Alignment(horizontal="center")
            )
            sheet.cell(row=excel_row, column=7, value=_safe(row["promo_reasons"])).alignment = (
                Alignment(wrap_text=True, vertical="top")
            )
            origin = origin_map.get(row["discovered_from_id"], "") if row["discovered_from_id"] else ""
            sheet.cell(row=excel_row, column=8, value=_safe(origin))
            sheet.cell(row=excel_row, column=9, value=_safe(local_str(row["discovered_at"])))
            sheet.cell(row=excel_row, column=10, value=_safe(local_str(row["evaluated_at"])))

    def _write_runs_sheet(self, sheet: Worksheet, rows: Sequence[sqlite3.Row]) -> None:
        _style_header(sheet, RUN_COLUMNS)
        for index, row in enumerate(rows, start=1):
            excel_row = index + 1
            duration = row["duration_ms"]
            sheet.cell(row=excel_row, column=1, value=_safe(local_str(row["started_at"])))
            sheet.cell(row=excel_row, column=2, value=_safe(local_str(row["finished_at"], "진행 중")))
            sheet.cell(
                row=excel_row,
                column=3,
                value=humanize_duration(duration / 1000 if duration else None),
            )
            for offset, key in enumerate(
                (
                    "sources_total",
                    "sources_ok",
                    "sources_failed",
                    "pages_fetched",
                    "candidates_found",
                    "new_sites",
                    "updated_sites",
                    "candidates_added",
                    "sources_approved",
                    "dropped_foreign",
                ),
                start=4,
            ):
                cell = sheet.cell(row=excel_row, column=offset, value=int(row[key] or 0))
                cell.alignment = Alignment(horizontal="center")
            sheet.cell(row=excel_row, column=14, value=_safe(row["note"]))

    def _write_summary_sheet(
        self,
        sheet: Worksheet,
        summary: dict[str, Any],
        total_rows: int,
        new_rows: int,
        contacts: int,
        url_count: int = 0,
    ) -> None:
        sheet.column_dimensions["A"].width = 26
        sheet.column_dimensions["B"].width = 30
        sheet.column_dimensions["C"].width = 14
        sheet.column_dimensions["D"].width = 12

        sheet["A1"] = "불법사이트 URL 수집 현황"
        sheet["A1"].font = TITLE_FONT
        sheet["A2"] = f"생성 시각: {local_str(now_utc())}"
        sheet["A2"].font = Font(color="808080", size=9)

        last_run = self.storage.last_run()
        rows: list[tuple[str, Any]] = [
            ("수집된 불법사이트 (URL 기준, 전체)", int(summary.get("total") or 0)),
            (
                f"엑셀 수록 ({GROUP_LABELS.get(self.group_mode, '사이트')} 기준, "
                f"{self.config.export.min_score}점 이상)",
                f"{total_rows}건  (URL {url_count}개를 묶음)",
            ),
            ("최근 24시간 신규", int(summary.get("new_24h") or 0)),
            ("최근 7일 신규", int(summary.get("new_7d") or 0)),
            ("생존 확인", int(summary.get("alive") or 0)),
            ("접속 불가", int(summary.get("dead") or 0)),
            ("생존 미확인", int(summary.get("unchecked") or 0)),
            ("신고 완료 처리", int(summary.get("reported") or 0)),
            ("제외 처리(오탐/홍보사이트 재분류)", int(summary.get("ignored") or 0)),
            ("연락 채널(텔레그램 등)", contacts),
            ("외국어로 제외한 도메인", int(summary.get("foreign_domains") or 0)),
            ("홍보사이트 등록/사용", f"{summary.get('sources_total', 0)} / {summary.get('sources_enabled', 0)}"),
            ("  그중 자동 발견", int(summary.get("sources_auto") or 0)),
            ("승인 대기 후보", int(summary.get("candidates_pending") or 0)),
            ("평가 대기 후보", int(summary.get("candidates_discovered") or 0)),
        ]
        if last_run is not None:
            duration = last_run["duration_ms"]
            rows.extend(
                [
                    ("마지막 사이클 시작", local_str(last_run["started_at"])),
                    ("마지막 사이클 종료", local_str(last_run["finished_at"], "진행 중")),
                    ("마지막 사이클 소요", humanize_duration(duration / 1000 if duration else None)),
                ]
            )

        row_index = 4
        for label, value in rows:
            label_cell = sheet.cell(row=row_index, column=1, value=label)
            label_cell.font = LABEL_FONT
            label_cell.border = THIN_BORDER
            value_cell = sheet.cell(row=row_index, column=2, value=value)
            value_cell.border = THIN_BORDER
            row_index += 1

        row_index += 1
        sheet.cell(row=row_index, column=1, value="카테고리별 집계").font = TITLE_FONT
        row_index += 1
        for header_index, header in enumerate(("카테고리", "건수"), start=1):
            cell = sheet.cell(row=row_index, column=header_index, value=header)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center")
        row_index += 1
        for entry in summary.get("by_category") or []:
            sheet.cell(row=row_index, column=1, value=entry.get("category_label") or entry.get("category"))
            sheet.cell(row=row_index, column=2, value=int(entry.get("n") or 0)).alignment = (
                Alignment(horizontal="center")
            )
            row_index += 1

        row_index += 1
        note = sheet.cell(
            row=row_index,
            column=1,
            value=(
                "※ 이 파일은 매 사이클마다 다시 생성됩니다. 엑셀에 직접 적은 메모는 사라지므로 "
                "처리 상태/메모는 `python bot.py mark <URL> --status reported --memo \"...\"` 로 "
                "남기세요(그러면 엑셀에도 계속 표시됩니다)."
            ),
        )
        note.font = Font(color="808080", size=9)
        note.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index + 2, end_column=4)

    # -- 저장 --------------------------------------------------------------
    def _save_workbook(self, workbook: Workbook, target: Path) -> tuple[Path, str]:
        """임시 파일에 쓴 뒤 교체합니다. 엑셀이 파일을 잡고 있으면 다른 이름으로 저장."""
        target.parent.mkdir(parents=True, exist_ok=True)
        fd, temp_name = tempfile.mkstemp(dir=str(target.parent), prefix=".export-", suffix=".xlsx")
        os.close(fd)
        temp_path = Path(temp_name)
        try:
            workbook.save(temp_path)
        except Exception:
            temp_path.unlink(missing_ok=True)
            raise

        try:
            os.replace(temp_path, target)
            return target, ""
        except (PermissionError, OSError) as exc:
            # Windows 에서 엑셀로 파일을 열어둔 경우 교체가 막힙니다.
            fallback = target.with_name(
                f"{target.stem}_열려있어_임시저장_{datetime.now():%H%M%S}{target.suffix}"
            )
            try:
                os.replace(temp_path, fallback)
            except OSError:
                temp_path.unlink(missing_ok=True)
                raise
            warning = (
                f"{target.name} 이 열려 있어 덮어쓰지 못했습니다({exc.__class__.__name__}). "
                f"{fallback.name} 으로 저장했습니다. 엑셀을 닫고 다시 내보내세요."
            )
            log.warning(warning)
            return fallback, warning

    def _write_snapshot(self, source: Path) -> Path | None:
        if not self.config.export.keep_daily_snapshots:
            return None
        snapshot = source.with_name(f"{source.stem}_{datetime.now():%Y%m%d}{source.suffix}")
        try:
            shutil.copy2(source, snapshot)
        except OSError as exc:
            log.warning("일별 스냅샷 저장 실패: %s", exc)
            return None
        self._prune_snapshots(source)
        return snapshot

    def _prune_snapshots(self, source: Path) -> None:
        retention = self.config.export.snapshot_retention_days
        if retention <= 0:
            return
        cutoff = now_utc().timestamp() - retention * 86_400
        for path in source.parent.glob(f"{source.stem}_*{source.suffix}"):
            try:
                if path.stat().st_mtime < cutoff:
                    path.unlink()
                    log.debug("오래된 스냅샷 삭제: %s", path.name)
            except OSError:
                continue

    # -- CSV (외부 시스템 입력용) ------------------------------------------
    #: 뒤에서 다시 분류하는 시스템이 읽어갈 컬럼. URL 이 첫 컬럼입니다.
    CSV_COLUMNS = (
        "url",
        "score",
        "first_seen_at",
        "last_seen_at",
        "host",
        "domain",
        "category",
        "language",
        "promo_site_count",
        "promo_sites",
        "alive",
        "http_status",
        "last_checked_at",
        "matched_keywords",
        "reasons",
        "status",
    )

    def _csv_rows(
        self, rows: Sequence[sqlite3.Row], source_map: dict[int, list[str]]
    ) -> list[dict[str, Any]]:
        """DB 행을 CSV 한 줄씩으로 바꿉니다 (묶지 않고 URL 단위 그대로)."""
        alive_text = {1: "alive", 0: "dead", None: ""}
        output: list[dict[str, Any]] = []
        for row in rows:
            sources = source_map.get(int(row["id"]), [])
            output.append(
                {
                    "url": row["url"],
                    "score": int(row["score"] or 0),
                    # 기계가 읽는 피드라 저장된 UTC ISO-8601 을 그대로 넘깁니다.
                    # 로컬 시간으로 바꾸면 받는 쪽이 시간대를 알 수 없습니다.
                    "first_seen_at": row["first_seen_at"],
                    "last_seen_at": row["last_seen_at"],
                    "host": row["host"],
                    "domain": row["domain"],
                    "category": row["category"],
                    "language": row["language"],
                    "promo_site_count": int(row["distinct_sources"] or 0),
                    "promo_sites": " | ".join(sources),
                    "alive": alive_text.get(row["alive"], ""),
                    "http_status": row["http_status"] or "",
                    "last_checked_at": row["last_checked_at"] or "",
                    "matched_keywords": row["matched_keywords"],
                    "reasons": row["reasons"],
                    "status": row["status"],
                }
            )
        return output

    def _save_csv(self, target: Path, rows: Sequence[dict[str, Any]]) -> Path | None:
        """임시 파일에 쓴 뒤 교체합니다.

        읽는 쪽이 절반만 쓰인 파일을 집어가지 않도록 원자적으로 바꿉니다.
        한글 엑셀에서 바로 열리도록 UTF-8 BOM(utf-8-sig)을 붙입니다.
        """
        target.parent.mkdir(parents=True, exist_ok=True)
        fd, temp_name = tempfile.mkstemp(dir=str(target.parent), prefix=".csv-", suffix=".csv")
        os.close(fd)
        temp_path = Path(temp_name)
        try:
            with temp_path.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.DictWriter(handle, fieldnames=list(self.CSV_COLUMNS))
                writer.writeheader()
                writer.writerows(rows)
            os.replace(temp_path, target)
            return target
        except OSError as exc:
            temp_path.unlink(missing_ok=True)
            log.warning("CSV 저장 실패 %s: %s", target, exc)
            return None

    def export_csv(self) -> tuple[Path | None, Path | None, int]:
        """외부 시스템이 읽어갈 CSV 를 만듭니다. ``(전체, 신규, 행 수)``.

        엑셀과 달리 **묶지 않고 URL 단위 그대로** 내보냅니다. 뒤에서 다시
        분류하는 쪽이 정보를 잃지 않게 하기 위함이고, 묶고 싶으면 host/domain
        컬럼으로 직접 묶을 수 있습니다.
        """
        if not self.config.export.csv_enabled:
            return None, None, 0

        min_score = self.config.export.csv_min_score
        rows = self.storage.sites_for_export(
            min_score, categories_excluded=(CONTACT_CATEGORY,), include_ignored=False
        )
        new_rows = [
            row
            for row in self.storage.sites_first_seen_since(
                days_ago(1), min_score, include_ignored=False
            )
            if row["category"] != CONTACT_CATEGORY
        ]

        source_map = self.storage.source_urls_grouped()
        directory = self.config.export_dir
        full = self._save_csv(directory / "urls.csv", self._csv_rows(rows, source_map))
        recent = self._save_csv(
            directory / "urls_new.csv", self._csv_rows(new_rows, source_map)
        )
        # 저장에 실패했으면 행 수도 0 으로 보고합니다. 안 그러면 파일이 없는데
        # "N행 저장" 이라고 알려주게 됩니다.
        written = len(rows) if full else 0
        log.info(
            "CSV 저장: 전체 %d행 / 최근 24시간 신규 %d행",
            written,
            len(new_rows) if recent else 0,
        )
        return full, recent, written

    # -- 공개 API ----------------------------------------------------------
    def export(self) -> ExportResult:
        """현재 DB 내용을 엑셀 파일로 씁니다."""
        min_score = self.config.export.min_score

        # 점수 필터는 **묶은 뒤에** 겁니다. URL 단위로 먼저 자르면, 묶어서
        # 붙는 중복 발견 가산점(묶기를 도입한 이유)이 기준선을 넘길 기회가
        # 사라집니다. 26점짜리 두 URL 이 묶여 32점이 되는 경우가 그렇습니다.
        sites = self.storage.sites_for_export(
            0, categories_excluded=(CONTACT_CATEGORY,), include_ignored=False
        )
        new_sites = [
            row
            for row in self.storage.sites_first_seen_since(
                days_ago(1), 0, include_ignored=False
            )
            if row["category"] != CONTACT_CATEGORY
        ]
        contacts = self.storage.sites_by_category(CONTACT_CATEGORY)
        summary = self.storage.summary(min_score)

        # 관측 테이블은 가장 큰 테이블이라 한 번만 읽어서 돌려 씁니다.
        source_map = self.storage.source_urls_grouped()

        site_groups = [
            group for group in self._group(sites, source_map) if group.score >= min_score
        ]
        new_groups = [
            group for group in self._group(new_sites, source_map) if group.score >= min_score
        ]
        # 연락 채널은 묶지 않습니다. t.me/계정A 와 t.me/계정B 는 호스트가 같아도
        # 서로 다른 채널이라 묶으면 한 줄로 뭉개집니다.
        contact_groups = self._group(contacts, source_map, mode="url")

        workbook = Workbook()
        summary_sheet = workbook.active
        summary_sheet.title = "요약"
        self._write_summary_sheet(
            summary_sheet, summary, len(site_groups), len(new_groups), len(contact_groups),
            url_count=sum(group.url_count for group in site_groups),
        )

        self._write_sites_sheet(workbook.create_sheet("불법사이트목록"), site_groups)
        self._write_sites_sheet(workbook.create_sheet("신규_최근24시간"), new_groups)
        self._write_sites_sheet(workbook.create_sheet("연락채널"), contact_groups, mode="url")
        all_sources = self.storage.source_rows()
        approved = [row for row in all_sources if row["state"] == "approved"]
        candidates = [
            row
            for row in all_sources
            if row["state"] in {"discovered", "pending", "rejected", "auto_disabled"}
        ]
        candidates.sort(
            key=lambda row: (row["state"] != "pending", -int(row["promo_score"] or 0))
        )
        origin_map = {int(row["id"]): row["url"] for row in all_sources}

        self._write_sources_sheet(workbook.create_sheet("홍보사이트_수집원"), approved)
        self._write_candidates_sheet(
            workbook.create_sheet("홍보사이트_후보"), candidates, origin_map
        )
        self._write_runs_sheet(workbook.create_sheet("실행이력"), self.storage.recent_runs(100))

        target = self.config.export_path
        saved_path, warning = self._save_workbook(workbook, target)
        snapshot = self._write_snapshot(saved_path) if not warning else None
        csv_path, csv_new_path, csv_rows = self.export_csv()

        log.info(
            "엑셀 저장 완료: %s (%s 기준 %d건 / URL %d건 / 신규 %d건 / 연락채널 %d건)",
            saved_path,
            GROUP_LABELS.get(self.group_mode, self.group_mode),
            len(site_groups),
            len(sites),
            len(new_groups),
            len(contact_groups),
        )
        return ExportResult(
            path=saved_path,
            snapshot=snapshot,
            rows=len(site_groups),
            new_rows=len(new_groups),
            contacts=len(contact_groups),
            url_rows=len(sites),
            csv_path=csv_path,
            csv_new_path=csv_new_path,
            csv_rows=csv_rows,
            warning=warning,
        )
